"""Export leads to CSV or Excel (.xlsx). Uses openpyxl for Excel."""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font

EXPORT_FIELDS = [
    ("name", "Business Name"),
    ("category", "Category"),
    ("phone", "Phone"),
    ("email", "Email"),
    ("website", "Website"),
    ("address", "Address"),
    ("city", "City"),
    ("state", "State"),
    ("country", "Country"),
    ("gmaps_url", "Google Maps"),
    ("review_count", "Reviews"),
    ("rating", "Rating"),
    ("facebook", "Facebook"),
    ("instagram", "Instagram"),
    ("linkedin", "LinkedIn"),
    ("twitter", "Twitter/X"),
    ("owner", "Owner/Operator"),
    ("has_website", "Has Website"),
    ("https", "HTTPS"),
    ("mobile_friendly", "Mobile Friendly"),
    ("website_score", "Website Score"),
    ("seo_score", "SEO Score"),
    ("modernization_score", "Modernization Score"),
    ("opportunity", "Opportunity"),
    ("recommendation", "Recommendation"),
]


def _cell(lead: dict, key: str):
    val = lead.get(key)
    if key in ("has_website", "https", "mobile_friendly"):
        return "Yes" if val else "No"
    return "" if val is None else val


def to_csv(leads: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in EXPORT_FIELDS])
    for lead in leads:
        writer.writerow([_cell(lead, key) for key, _ in EXPORT_FIELDS])
    return buf.getvalue().encode("utf-8")


def to_xlsx(leads: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([label for _, label in EXPORT_FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for lead in leads:
        ws.append([_cell(lead, key) for key, _ in EXPORT_FIELDS])
    # Reasonable column widths.
    for i, (_, label) in enumerate(EXPORT_FIELDS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(label) + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
